"""
Driver SMS Desktop Application - Modern UI
A beautiful Windows desktop app for managing driver SMS communications via Twilio
Built with CustomTkinter for modern appearance
"""

import customtkinter as ctk
from tkinter import filedialog, messagebox
import json
import os
from datetime import datetime
import threading
import time
import requests
from requests.auth import HTTPBasicAuth
import openpyxl

# Set appearance mode and color theme
ctk.set_appearance_mode("dark")  # "dark" or "light"
ctk.set_default_color_theme("blue")

class ModernDriverSMSApp:
    def __init__(self):
        self.root = ctk.CTk()
        self.root.title("Driver SMS Manager")
        self.root.geometry("1400x800")
        
        # Data storage
        self.drivers = []
        self.conversations = {}
        self.selected_driver = None
        self.credentials = self.load_credentials()
        self.sort_by = "name"  # or "team"
        self.sort_reverse = False
        
        # Setup UI
        self.setup_ui()
        self.load_drivers_from_file()
        self.update_driver_list()
        
        # Start message polling thread
        self.polling = True
        self.poll_thread = threading.Thread(target=self.poll_messages, daemon=True)
        self.poll_thread.start()
        
        self.root.mainloop()
    
    def setup_ui(self):
        # Top navbar
        navbar = ctk.CTkFrame(self.root, height=70, corner_radius=0, fg_color=("#1e40af", "#1e3a8a"))
        navbar.pack(side="top", fill="x")
        navbar.pack_propagate(False)
        
        title_frame = ctk.CTkFrame(navbar, fg_color="transparent")
        title_frame.pack(side="left", padx=20, pady=15)
        
        ctk.CTkLabel(title_frame, text="💬 Driver SMS", font=ctk.CTkFont(size=24, weight="bold"),
                    text_color="white").pack()
        
        # Right side buttons
        button_frame = ctk.CTkFrame(navbar, fg_color="transparent")
        button_frame.pack(side="right", padx=20, pady=15)
        
        ctk.CTkButton(button_frame, text="📁 Load Excel", command=self.load_excel,
                     width=120, height=40, corner_radius=8, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        
        ctk.CTkButton(button_frame, text="📤 Bulk Send", command=self.show_bulk_message,
                     width=120, height=40, corner_radius=8, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        
        ctk.CTkButton(button_frame, text="+ Add Contact", command=self.add_contact,
                     width=120, height=40, corner_radius=8, font=ctk.CTkFont(size=13)).pack(side="left", padx=5)
        
        ctk.CTkButton(button_frame, text="⚙️ Settings", command=self.show_settings,
                     width=100, height=40, corner_radius=8, font=ctk.CTkFont(size=13),
                     fg_color=("#374151", "#1f2937")).pack(side="left", padx=5)
        
        # Main container
        main_container = ctk.CTkFrame(self.root, fg_color="transparent")
        main_container.pack(fill="both", expand=True, padx=10, pady=10)
        
        # Left panel - Driver list
        left_panel = ctk.CTkFrame(main_container, width=400, corner_radius=10)
        left_panel.pack(side="left", fill="both", padx=(0, 5))
        left_panel.pack_propagate(False)
        
        # Search bar
        search_frame = ctk.CTkFrame(left_panel, fg_color="transparent")
        search_frame.pack(fill="x", padx=15, pady=15)
        
        self.search_var = ctk.StringVar()
        # Remove the auto-refresh on every keystroke
        
        self.search_entry = ctk.CTkEntry(search_frame, placeholder_text="🔍 Search drivers...",
                                        textvariable=self.search_var, height=40,
                                        font=ctk.CTkFont(size=13))
        self.search_entry.pack(fill="x")
        self.search_entry.bind('<Return>', lambda e: self.update_driver_list())  # Only search on Enter
        
        # Sort buttons
        sort_frame = ctk.CTkFrame(left_panel, fg_color="transparent")
        sort_frame.pack(fill="x", padx=15, pady=(0, 10))
        
        ctk.CTkLabel(sort_frame, text="Sort by:", font=ctk.CTkFont(size=12)).pack(side="left", padx=(0, 10))
        
        ctk.CTkButton(sort_frame, text="Name", command=lambda: self.set_sort("name"),
                     width=80, height=30, corner_radius=6, font=ctk.CTkFont(size=11)).pack(side="left", padx=2)
        
        ctk.CTkButton(sort_frame, text="Team", command=lambda: self.set_sort("team"),
                     width=80, height=30, corner_radius=6, font=ctk.CTkFont(size=11)).pack(side="left", padx=2)
        
        # Driver list with scrollbar
        list_frame = ctk.CTkScrollableFrame(left_panel, corner_radius=10)
        list_frame.pack(fill="both", expand=True, padx=15, pady=(0, 15))
        
        self.driver_buttons = []
        self.driver_list_frame = list_frame
        
        # Right panel - Chat
        right_panel = ctk.CTkFrame(main_container, corner_radius=10)
        right_panel.pack(side="left", fill="both", expand=True)
        
        # Chat header
        chat_header = ctk.CTkFrame(right_panel, height=60, corner_radius=0, fg_color=("#e5e7eb", "#374151"))
        chat_header.pack(fill="x")
        chat_header.pack_propagate(False)
        
        self.header_label = ctk.CTkLabel(chat_header, text="Select a conversation",
                                        font=ctk.CTkFont(size=16, weight="bold"))
        self.header_label.pack(side="left", padx=20, pady=15)
        
        # Messages area
        self.messages_frame = ctk.CTkTextbox(right_panel, corner_radius=0, font=ctk.CTkFont(size=13))
        self.messages_frame.pack(fill="both", expand=True, padx=0, pady=0)
        
        # Message input
        input_container = ctk.CTkFrame(right_panel, fg_color="transparent")
        input_container.pack(fill="x", padx=15, pady=15)
        
        self.message_entry = ctk.CTkEntry(input_container, placeholder_text="Type a message...",
                                         height=50, font=ctk.CTkFont(size=14))
        self.message_entry.pack(side="left", fill="both", expand=True, padx=(0, 10))
        self.message_entry.bind('<Return>', lambda e: self.send_message())
        
        ctk.CTkButton(input_container, text="Send", command=self.send_message,
                     width=100, height=50, corner_radius=8, font=ctk.CTkFont(size=14, weight="bold")).pack(side="right")
    
    def set_sort(self, sort_by):
        """Set sorting method"""
        if self.sort_by == sort_by:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_by = sort_by
            self.sort_reverse = False
        self.update_driver_list()
    
    def update_driver_list(self):
        """Update the driver list with current sort and filter"""
        # Clear existing buttons
        for widget in self.driver_list_frame.winfo_children():
            widget.destroy()
        self.driver_buttons.clear()
        
        # Filter drivers
        search = self.search_var.get().lower()
        filtered = [d for d in self.drivers if search in d['name'].lower() or search in d.get('team', '').lower()]
        
        # Sort drivers
        if self.sort_by == "name":
            filtered.sort(key=lambda x: x['name'], reverse=self.sort_reverse)
        else:
            filtered.sort(key=lambda x: x.get('team', ''), reverse=self.sort_reverse)
        
        # Create driver buttons
        for driver in filtered:
            has_messages = driver['phone'] in self.conversations
            
            driver_frame = ctk.CTkFrame(self.driver_list_frame, corner_radius=8,
                                       fg_color=("#f3f4f6", "#2d3748"))
            driver_frame.pack(fill="x", pady=5)
            
            # Create button with proper text layout
            display_text = f"{driver['name']}\n{driver.get('team', 'No team')}"
            if has_messages:
                display_text += " 💬"
            
            btn = ctk.CTkButton(driver_frame, text=display_text, corner_radius=8, height=70,
                              fg_color="transparent", hover_color=("#e5e7eb", "#374151"),
                              anchor="w", command=lambda d=driver: self.select_driver(d),
                              font=ctk.CTkFont(size=13))
            btn.pack(fill="both", expand=True, padx=5, pady=5)
            
            self.driver_buttons.append((btn, driver))
    
    def select_driver(self, driver):
        """Select a driver to view conversation"""
        self.selected_driver = driver
        self.display_conversation()
    
    def display_conversation(self):
        """Display conversation with selected driver"""
        if not self.selected_driver:
            return
        
        self.header_label.configure(text=f"{self.selected_driver['name']} - {self.selected_driver['phone']}")
        
        self.messages_frame.delete("1.0", "end")
        
        messages = self.conversations.get(self.selected_driver['phone'], [])
        messages.sort(key=lambda x: x['timestamp'])
        
        for msg in messages:
            if msg['direction'] == 'outbound-api':
                self.messages_frame.insert("end", f"You: {msg['body']}\n\n")
            else:
                self.messages_frame.insert("end", f"{self.selected_driver['firstName']}: {msg['body']}\n\n")
        
        self.messages_frame.see("end")
    
    def load_credentials(self):
        """Load saved Twilio credentials"""
        if os.path.exists("config.json"):
            with open("config.json", 'r') as f:
                return json.load(f)
        return {"account_sid": "", "auth_token": "", "twilio_number": ""}
    
    def save_credentials(self):
        """Save Twilio credentials"""
        with open("config.json", 'w') as f:
            json.dump(self.credentials, f)
    
    def load_drivers_from_file(self):
        """Load saved drivers"""
        if os.path.exists("drivers.json"):
            with open("drivers.json", 'r') as f:
                self.drivers = json.load(f)
    
    def save_drivers(self):
        """Save drivers to file"""
        with open("drivers.json", 'w') as f:
            json.dump(self.drivers, f, indent=2)
    
    def load_excel(self):
        """Load drivers from Excel file"""
        filename = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        
        if not filename:
            return
        
        try:
            wb = openpyxl.load_workbook(filename)
            ws = wb.active
            
            new_drivers = []
            headers = [cell.value for cell in ws[1]]
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                
                if row_dict.get('Phone') and row_dict.get('Active') == 'Yes':
                    driver = {
                        'name': row_dict.get('Full Name', ''),
                        'phone': row_dict.get('Phone', ''),
                        'team': row_dict.get('Active Team', ''),
                        'firstName': row_dict.get('Full Name', '').split()[0] if row_dict.get('Full Name') else ''
                    }
                    new_drivers.append(driver)
            
            self.drivers.extend(new_drivers)
            self.save_drivers()
            self.update_driver_list()
            
            messagebox.showinfo("Success", f"Loaded {len(new_drivers)} drivers!")
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load Excel: {str(e)}")
    
    def add_contact(self):
        """Add a single contact manually"""
        dialog = ctk.CTkToplevel(self.root)
        dialog.title("Add Contact")
        dialog.geometry("450x300")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ctk.CTkLabel(dialog, text="Name:", font=ctk.CTkFont(size=14)).pack(pady=(30, 5))
        name_entry = ctk.CTkEntry(dialog, width=350, height=40, font=ctk.CTkFont(size=13))
        name_entry.pack(pady=5)
        
        ctk.CTkLabel(dialog, text="Phone Number:", font=ctk.CTkFont(size=14)).pack(pady=(15, 5))
        phone_entry = ctk.CTkEntry(dialog, width=350, height=40, font=ctk.CTkFont(size=13))
        phone_entry.pack(pady=5)
        phone_entry.insert(0, "+1")
        
        def save_contact():
            name = name_entry.get().strip()
            phone = phone_entry.get().strip()
            
            if not name or not phone:
                messagebox.showwarning("Error", "Please fill in all fields")
                return
            
            driver = {
                'name': name,
                'phone': phone,
                'team': 'Manual',
                'firstName': name.split()[0]
            }
            
            self.drivers.append(driver)
            self.save_drivers()
            self.update_driver_list()
            dialog.destroy()
            messagebox.showinfo("Success", f"Added {name}!")
        
        ctk.CTkButton(dialog, text="Add Contact", command=save_contact,
                     width=200, height=45, corner_radius=8, 
                     font=ctk.CTkFont(size=14, weight="bold")).pack(pady=30)
    
    def show_settings(self):
        """Show settings dialog for Twilio credentials"""
        dialog = ctk.CTkToplevel(self.root)
        dialog.title("Twilio Settings")
        dialog.geometry("550x400")
        dialog.transient(self.root)
        dialog.grab_set()
        
        ctk.CTkLabel(dialog, text="Account SID:", font=ctk.CTkFont(size=14)).pack(pady=(30, 5))
        sid_entry = ctk.CTkEntry(dialog, width=450, height=40, font=ctk.CTkFont(size=13))
        sid_entry.pack(pady=5)
        sid_entry.insert(0, self.credentials.get('account_sid', ''))
        
        ctk.CTkLabel(dialog, text="Auth Token:", font=ctk.CTkFont(size=14)).pack(pady=(15, 5))
        token_entry = ctk.CTkEntry(dialog, width=450, height=40, show="*", font=ctk.CTkFont(size=13))
        token_entry.pack(pady=5)
        token_entry.insert(0, self.credentials.get('auth_token', ''))
        
        ctk.CTkLabel(dialog, text="Twilio Phone Number:", font=ctk.CTkFont(size=14)).pack(pady=(15, 5))
        phone_entry = ctk.CTkEntry(dialog, width=450, height=40, font=ctk.CTkFont(size=13))
        phone_entry.pack(pady=5)
        phone_entry.insert(0, self.credentials.get('twilio_number', ''))
        
        def save_settings():
            self.credentials['account_sid'] = sid_entry.get().strip()
            self.credentials['auth_token'] = token_entry.get().strip()
            self.credentials['twilio_number'] = phone_entry.get().strip()
            self.save_credentials()
            dialog.destroy()
            messagebox.showinfo("Success", "Settings saved!")
        
        ctk.CTkButton(dialog, text="Save Settings", command=save_settings,
                     width=250, height=50, corner_radius=8,
                     font=ctk.CTkFont(size=15, weight="bold")).pack(pady=30)
    
    def show_bulk_message(self):
        """Show bulk message dialog"""
        if not self.credentials.get('account_sid'):
            messagebox.showwarning("Setup Required", "Please configure Twilio settings first!")
            self.show_settings()
            return
        
        if not self.drivers:
            messagebox.showwarning("No Drivers", "Please load drivers first!")
            return
        
        dialog = ctk.CTkToplevel(self.root)
        dialog.title("Send Bulk Message")
        dialog.geometry("700x650")
        dialog.transient(self.root)
        dialog.grab_set()
        
        teams = sorted(set(d.get('team', '') for d in self.drivers if d.get('team')))
        
        ctk.CTkLabel(dialog, text="Select Team:", font=ctk.CTkFont(size=15, weight="bold")).pack(pady=(20, 5))
        
        team_var = ctk.StringVar(value="All Teams")
        team_menu = ctk.CTkOptionMenu(dialog, variable=team_var, values=["All Teams"] + teams,
                                      width=600, height=40, font=ctk.CTkFont(size=13))
        team_menu.pack(pady=10)
        
        ctk.CTkLabel(dialog, text="Message:", font=ctk.CTkFont(size=15, weight="bold")).pack(pady=(15, 5))
        
        message_text = ctk.CTkTextbox(dialog, width=600, height=150, font=ctk.CTkFont(size=13))
        message_text.pack(pady=10)
        message_text.insert("1.0", "Hi {firstName}, ")
        
        ctk.CTkLabel(dialog, text="Use: {firstName}, {name}, {team}", 
                    font=ctk.CTkFont(size=11), text_color="gray").pack()
        
        selected_count = ctk.CTkLabel(dialog, text="", font=ctk.CTkFont(size=13))
        selected_count.pack(pady=10)
        
        def update_count(*args):
            team = team_var.get()
            count = len(self.drivers) if team == "All Teams" else len([d for d in self.drivers if d.get('team') == team])
            cost = count * 0.0079
            selected_count.configure(text=f"Will send to {count} drivers (Est. cost: ${cost:.2f})")
        
        team_var.trace('w', update_count)
        update_count()
        
        def send_bulk():
            message = message_text.get("1.0", "end").strip()
            if not message:
                messagebox.showwarning("No Message", "Please enter a message")
                return
            
            team = team_var.get()
            drivers = self.drivers if team == "All Teams" else [d for d in self.drivers if d.get('team') == team]
            
            if not messagebox.askyesno("Confirm", f"Send to {len(drivers)} driver(s)?"):
                return
            
            dialog.destroy()
            self.send_bulk_messages(drivers, message)
        
        ctk.CTkButton(dialog, text="Send Messages", command=send_bulk,
                     width=300, height=50, corner_radius=8,
                     font=ctk.CTkFont(size=16, weight="bold")).pack(pady=20)
    
    def send_bulk_messages(self, drivers, message_template):
        """Send bulk messages to selected drivers"""
        success = 0
        failed = 0
        
        for driver in drivers:
            msg = message_template.replace('{firstName}', driver.get('firstName', ''))
            msg = msg.replace('{name}', driver['name'])
            msg = msg.replace('{team}', driver.get('team', ''))
            
            if self.send_sms(driver['phone'], msg):
                success += 1
            else:
                failed += 1
            
            time.sleep(0.1)
        
        messagebox.showinfo("Complete", f"Sent: {success}\nFailed: {failed}")
    
    def send_sms(self, to_number, message):
        """Send SMS via Twilio"""
        try:
            url = f"https://api.twilio.com/2010-04-01/Accounts/{self.credentials['account_sid']}/Messages.json"
            
            response = requests.post(
                url,
                auth=HTTPBasicAuth(self.credentials['account_sid'], self.credentials['auth_token']),
                data={
                    'To': to_number,
                    'From': self.credentials['twilio_number'],
                    'Body': message
                }
            )
            
            return response.status_code == 201
        except Exception as e:
            print(f"Error sending SMS: {e}")
            return False
    
    def send_message(self):
        """Send message to selected driver"""
        if not self.selected_driver:
            messagebox.showwarning("No Selection", "Please select a driver first")
            return
        
        message = self.message_entry.get().strip()
        if not message:
            return
        
        if self.send_sms(self.selected_driver['phone'], message):
            self.message_entry.delete(0, "end")
            
            if self.selected_driver['phone'] not in self.conversations:
                self.conversations[self.selected_driver['phone']] = []
            
            self.conversations[self.selected_driver['phone']].append({
                'sid': 'local_' + str(time.time()),
                'body': message,
                'direction': 'outbound-api',
                'timestamp': datetime.now().isoformat()
            })
            
            self.display_conversation()
        else:
            messagebox.showerror("Error", "Failed to send message")
    
    def poll_messages(self):
        """Poll Twilio for new messages"""
        while self.polling:
            if self.credentials.get('account_sid'):
                try:
                    self.fetch_messages()
                except Exception as e:
                    print(f"Polling error: {e}")
            time.sleep(10)
    
    def fetch_messages(self):
        """Fetch messages from Twilio"""
        url = f"https://api.twilio.com/2010-04-01/Accounts/{self.credentials['account_sid']}/Messages.json"
        
        try:
            response = requests.get(
                url,
                auth=HTTPBasicAuth(self.credentials['account_sid'], self.credentials['auth_token']),
                params={'PageSize': 50}
            )
            
            if response.status_code == 200:
                messages = response.json().get('messages', [])
                for msg in messages:
                    phone = msg['from'] if msg['direction'] == 'inbound' else msg['to']
                    
                    if phone not in self.conversations:
                        self.conversations[phone] = []
                    
                    if not any(m['sid'] == msg['sid'] for m in self.conversations[phone]):
                        self.conversations[phone].append({
                            'sid': msg['sid'],
                            'body': msg['body'],
                            'direction': msg['direction'],
                            'timestamp': msg['date_created']
                        })
                
                self.update_driver_list()
                
        except Exception as e:
            print(f"Error fetching messages: {e}")

def main():
    app = ModernDriverSMSApp()

if __name__ == "__main__":
    main()
