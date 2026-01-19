import markdown
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# Read the markdown file
with open('SYSTEM_OVERVIEW.md', 'r', encoding='utf-8') as f:
    md_content = f.read()

# Create a new Word document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Process markdown line by line
lines = md_content.split('\n')
i = 0

while i < len(lines):
    line = lines[i].rstrip()
    
    # Skip empty lines
    if not line:
        i += 1
        continue
    
    # Heading 1
    if line.startswith('# '):
        p = doc.add_heading(line[2:], level=1)
        p.runs[0].font.color.rgb = RGBColor(79, 70, 229)  # Indigo color
    
    # Heading 2
    elif line.startswith('## '):
        p = doc.add_heading(line[3:], level=2)
        p.runs[0].font.color.rgb = RGBColor(99, 102, 241)  # Lighter indigo
    
    # Heading 3
    elif line.startswith('### '):
        p = doc.add_heading(line[4:], level=3)
    
    # Code blocks
    elif line.startswith('```'):
        # Find the end of code block
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code_lines.append(lines[i])
            i += 1
        
        # Add code block with gray background
        p = doc.add_paragraph()
        p.style = 'No Spacing'
        run = p.add_run('\n'.join(code_lines))
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        
        # Add shading (light gray background)
        from docx.oxml.shared import OxmlElement
        from docx.oxml.ns import qn
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), 'F3F4F6')
        p._element.get_or_add_pPr().append(shading_elm)
    
    # Bullet lists
    elif line.startswith('- '):
        text = line[2:]
        # Remove markdown formatting
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # Bold
        text = re.sub(r'`(.*?)`', r'\1', text)  # Code
        text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)  # Links
        doc.add_paragraph(text, style='List Bullet')
    
    # Checkbox lists
    elif line.startswith('- [ ] ') or line.startswith('- [x] ') or line.startswith('- [✅] ') or line.startswith('- [⚠️] '):
        text = line[6:] if line.startswith('- [ ] ') else line[7:]
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        p = doc.add_paragraph(text, style='List Bullet')
    
    # Horizontal rule
    elif line.startswith('---'):
        p = doc.add_paragraph()
        p.add_run('_' * 80)
    
    # Regular paragraphs
    else:
        # Remove markdown formatting
        text = re.sub(r'\*\*(.*?)\*\*', lambda m: m.group(1), line)  # Bold
        text = re.sub(r'`(.*?)`', lambda m: m.group(1), text)  # Inline code
        text = re.sub(r'\[(.*?)\]\((.*?)\)', lambda m: f'{m.group(1)} ({m.group(2)})', text)  # Links
        
        if text.strip():
            p = doc.add_paragraph(text)
    
    i += 1

# Save the document
doc.save('SYSTEM_OVERVIEW.docx')
print("Successfully converted SYSTEM_OVERVIEW.md to SYSTEM_OVERVIEW.docx")
